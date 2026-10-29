"""Læser og skriver tabellerne i K-flow.xlsx (Ordrer, Ordrelinjer, Kanpla, Kolonneregler).

I drift skriver Office Scriptet ordrerne. Her bruges det til:
  * at læse arbejdsbogen og lave data til dashboardet
  * at indlæse gamle mails eller test-mails (samme logik som scriptet: dubletter springes over)
  * at indlæse Kanpla-tal
"""
from datetime import date, datetime, timedelta

from openpyxl.utils import get_column_letter as L
from openpyxl.utils.cell import range_boundaries

KR = '#,##0.00;-#,##0.00;"-"'
EXCEL_EPOKE = date(1899, 12, 30)


def _find_tabel(wb, navn):
    for ws in wb.worksheets:
        if navn in ws.tables:
            return ws, ws.tables[navn]
    raise KeyError(f"Tabellen '{navn}' findes ikke i arbejdsbogen")


def _som_dato(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and v > 0:      # Excel-serienummer (skrevet af Office Scriptet)
        return EXCEL_EPOKE + timedelta(days=int(v))
    if isinstance(v, str) and v:
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
            try:
                return datetime.strptime(v, fmt).date()
            except ValueError:
                pass
    return None


def _tekst(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return "" if v is None else str(v)


def laes_tabel(wb, navn):
    ws, t = _find_tabel(wb, navn)
    c1, r1, c2, r2 = range_boundaries(t.ref)
    rows = list(ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2, values_only=True))
    hoved = [_tekst(h) for h in rows[0]]
    return [dict(zip(hoved, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]


def tilfoej(wb, navn, raekker, datokolonner=("Leveringsdato", "Dato"), kr_kolonner=()):
    """Tilføjer rækker (dicts) nederst i tabellen og udvider tabellens område."""
    if not raekker:
        return
    ws, t = _find_tabel(wb, navn)
    c1, r1, c2, r2 = range_boundaries(t.ref)
    hoved = [_tekst(ws.cell(r1, c).value) for c in range(c1, c2 + 1)]
    # Tom skabelon-række fyldes først
    r = r2 if r2 == r1 + 1 and all(ws.cell(r2, c).value in (None, "") for c in range(c1, c2 + 1)) else r2 + 1
    for raekke in raekker:
        for i, h in enumerate(hoved):
            celle = ws.cell(r, c1 + i, raekke.get(h))
            if h in datokolonner:
                celle.number_format = "dd-mm-yyyy"
            elif h in kr_kolonner:
                celle.number_format = KR
        r += 1
    t.ref = f"{L(c1)}{r1}:{L(c2)}{r - 1}"


def laes_regler(wb):
    return [{"leverandoer": _tekst(r.get("Leverandør")), "genkend": _tekst(r.get("Genkend i mail")),
             "kunde_indeholder": _tekst(r.get("Kunde indeholder")), "kolonne": _tekst(r.get("Kolonne"))}
            for r in laes_tabel(wb, "Kolonneregler")]


def importer_ordre(wb, o):
    """Samme som Office Scriptet: returnerer 'indlæst', 'dublet', 'ukendt leverandør' eller 'fejl'."""
    if not o["ordre"] or not o["linjer"]:
        return "fejl"
    if any(_tekst(r["Ordre"]) == o["ordre"] for r in laes_tabel(wb, "Ordrer")):
        return "dublet"
    dato = _som_dato(o["leveringsdato"])
    faelles = {"Ordre": o["ordre"], "Leveringsdato": dato, "Leverandør": o["leverandoer"], "Kunde": o["kunde"],
               "Bruger": o["bruger"], "Leveringsbemærkning": o["leveringsbemaerkning"]}
    tilfoej(wb, "Ordrer", [{**faelles, "Kolonne": o["kolonne"], "Antal linjer": len(o["linjer"]),
                            "Total": o["total"], "Kilde": o["kilde"],
                            "Indlæst": datetime.now().isoformat(timespec="minutes"),
                            "Advarsler": "; ".join(o["advarsler"])}], kr_kolonner=("Total",))
    tilfoej(wb, "Ordrelinjer", [{**faelles, "Varenr": l["varenr"], "Vare": l["vare"], "Antal": l["antal"],
                                 "Enhed": l["enhed"], "Enhedspris": l["enhedspris"], "Linjepris": l["linjepris"]}
                                for l in o["linjer"]], kr_kolonner=("Enhedspris", "Linjepris"))
    return "indlæst" if o["leverandoer"] != "Ukendt" else "ukendt leverandør"


def importer_kanpla(wb, dage):
    """Tilføjer Kanpla-dage, som ikke allerede står i tabellen. Returnerer antal nye dage."""
    findes = {_som_dato(r["Dato"]) for r in laes_tabel(wb, "Kanpla")}
    nye = [{"Dato": date.fromisoformat(d), **v} for d, v in sorted(dage.items())
           if date.fromisoformat(d) not in findes]
    tilfoej(wb, "Kanpla", nye)
    return len(nye)


def laes_data(wb):
    """Ordrer (med varelinjer) og Kanpla-dage i samme form, som beregning.py og dashboardet bruger."""
    linjer = {}
    for r in laes_tabel(wb, "Ordrelinjer"):
        linjer.setdefault(_tekst(r["Ordre"]), []).append({
            "varenr": _tekst(r["Varenr"]), "vare": _tekst(r["Vare"]), "antal": r["Antal"] or 0,
            "enhed": _tekst(r["Enhed"]), "enhedspris": r["Enhedspris"] or 0, "linjepris": r["Linjepris"] or 0})
    ordrer = []
    for r in laes_tabel(wb, "Ordrer"):
        d = _som_dato(r["Leveringsdato"])
        nr = _tekst(r["Ordre"])
        ordrer.append({
            "ordre": nr, "leveringsdato": d.isoformat() if d else "", "leverandoer": _tekst(r["Leverandør"]),
            "kunde": _tekst(r["Kunde"]), "bruger": _tekst(r["Bruger"]),
            "leveringsbemaerkning": _tekst(r["Leveringsbemærkning"]), "kolonne": _tekst(r["Kolonne"]),
            "total": r["Total"] or 0, "kilde": _tekst(r["Kilde"]), "indlaest": _tekst(r["Indlæst"]),
            "advarsler": [a for a in _tekst(r["Advarsler"]).split("; ") if a],
            "linjer": linjer.get(nr, [])})
    kanpla = {}
    for r in laes_tabel(wb, "Kanpla"):
        d = _som_dato(r["Dato"])
        if d:
            kanpla[d.isoformat()] = {k: (v or 0) for k, v in r.items() if k != "Dato"}
    return ordrer, kanpla
