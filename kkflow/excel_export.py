"""Bygger Excel-filen. Månedsarket ligner det manuelle ark, men henter tallene med formler fra
arkene 'Ordrer' og 'Kanpla', så det kan genberegnes og efterprøves direkte i Excel."""
import calendar
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo

GUL = PatternFill("solid", fgColor="FFFFCC")
BLAA = PatternFill("solid", fgColor="DDEBF7")   # udfyldes automatisk
GRAA = PatternFill("solid", fgColor="BFBFBF")
FED = Font(bold=True)
TYND = Side(style="thin", color="999999")
RAMME = Border(left=TYND, right=TYND, top=TYND, bottom=TYND)
KR = '#,##0.00;-#,##0.00;"-"'
MDR = ["januar", "februar", "marts", "april", "maj", "juni", "juli", "august",
       "september", "oktober", "november", "december"]


def _tabel(ws, navn, overskrifter, raekker, bredder):
    ws.append(overskrifter)
    for r in raekker:
        ws.append(r)
    for i, b in enumerate(bredder, 1):
        ws.column_dimensions[L(i)].width = b
    ref = f"A1:{L(len(overskrifter))}{max(len(raekker) + 1, 2)}"
    t = Table(displayName=navn, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = "A2"


def _d(iso):
    return date.fromisoformat(iso) if iso else None


def ordrer_ark(wb, ordrer):
    ws = wb.create_sheet("Ordrer")
    _tabel(ws, "Ordrer",
           ["Ordre", "Leveringsdato", "Leverandør", "Kunde", "Bruger", "Leveringsbemærkning",
            "Kolonne", "Antal linjer", "Total", "Kildefil"],
           [[o["ordre"], _d(o["leveringsdato"]), o["leverandoer"], o["kunde"], o["bruger"],
             o["leveringsbemaerkning"], o["kolonne"], len(o["linjer"]), o["total"], o["kilde"]]
            for o in ordrer],
           [11, 13, 22, 38, 26, 22, 14, 11, 12, 30])
    for c in ws["B"][1:]:
        c.number_format = "dd-mm-yyyy"
    for c in ws["I"][1:]:
        c.number_format = KR


def linje_ark(wb, ordrer):
    ws = wb.create_sheet("Ordrelinjer")
    raekker = [[o["ordre"], _d(o["leveringsdato"]), o["leverandoer"], o["kunde"], o["bruger"],
                o["leveringsbemaerkning"], l["varenr"], l["vare"], l["antal"], l["enhed"],
                l["enhedspris"], l["linjepris"]]
               for o in ordrer for l in o["linjer"]]
    _tabel(ws, "Ordrelinjer",
           ["Ordre", "Leveringsdato", "Leverandør", "Kunde", "Bruger", "Leveringsbemærkning",
            "Varenr", "Vare", "Antal", "Enhed", "Enhedspris", "Linjepris"],
           raekker, [11, 13, 22, 38, 26, 22, 9, 30, 8, 8, 12, 12])
    for c in ws["B"][1:]:
        c.number_format = "dd-mm-yyyy"
    for col in ("K", "L"):
        for c in ws[col][1:]:
            c.number_format = KR


def kanpla_ark(wb, kanpla, cfg):
    ws = wb.create_sheet("Kanpla")
    noegler = ["kuverter"] + [k["noegle"] for k in cfg["ekstrasalg"]]
    _tabel(ws, "Kanpla", ["Dato"] + noegler,
           [[_d(d)] + [v.get(k, 0.0) for k in noegler] for d, v in sorted(kanpla.items())],
           [12] + [16] * len(noegler))
    for c in ws["A"][1:]:
        c.number_format = "dd-mm-yyyy"
    return {k: L(i + 2) for i, k in enumerate(noegler)}


def maaneds_ark(wb, aar, md, cfg, kanpla_kol):
    ws = wb.create_sheet(f"{MDR[md - 1].capitalize()} {aar}", 0)
    kol = cfg["indkoebs_kolonner"]
    eks = cfg["ekstrasalg"]
    n_dage = calendar.monthrange(aar, md)[1]
    F, S = 9, 9 + n_dage - 1           # første og sidste datorække
    lev0 = 2                           # B
    lev_sidst = lev0 + len(kol) - 1    # H
    c_ialt, c_mtd = lev_sidst + 1, lev_sidst + 2          # I, J
    c_oevr = c_mtd + 2                                    # L
    c_kp, c_kuv, c_kb, c_kpk, c_lb, c_diff = range(c_oevr + 2, c_oevr + 8)  # N..S
    e0 = c_diff + 2                                       # U
    c_eks, c_vare, c_tilv = e0 + len(eks), e0 + len(eks) + 1, e0 + len(eks) + 2

    ws["A2"], ws["B2"] = "Måned", f"{MDR[md - 1]} {aar}"
    ws["A3"], ws["B3"] = "Kantine", cfg["kantine"]
    ws["A4"], ws["C4"] = "Kuvertpris", cfg["kuvertpris"]
    for c in ("B2", "B3", "C4"):
        ws[c].fill = GUL
        ws[c].border = RAMME
    ws["C4"].number_format = '"kr." #,##0.00'
    ws["A5"] = "Gule felter: indtastes manuelt.  Blå felter: hentes automatisk fra mails (Ordrer) og Kanpla."
    ws["A5"].font = Font(italic=True, size=9)

    def hoved(c, tekst, r=8):
        cell = ws.cell(r, c, tekst)
        cell.font, cell.fill, cell.border = FED, GRAA, RAMME
        cell.alignment = Alignment(horizontal="center", vertical="bottom", wrap_text=True)

    hoved(1, "Dato")
    for i, k in enumerate(kol):
        hoved(lev0 + i, k["konto"], 7)
        hoved(lev0 + i, k["navn"])
    for c, t in [(c_ialt, "Indkøb i alt"), (c_mtd, "MTD"), (c_oevr, "Øvrige omkostninger"),
                 (c_kp, "Kuv pris"), (c_kuv, "Kuverter"), (c_kb, "Kuvert Budget"), (c_kpk, "Kuvertpris"),
                 (c_lb, "Løbende Budget"), (c_diff, "Diff fra Budget"),
                 (c_eks, "Ekstra salg"), (c_vare, "Vareforbrug"), (c_tilv, "Til vareforbrug")]:
        hoved(c, t)
    hoved(c_oevr, "Ex moms", 7)
    hoved(e0, "Ekstrasalg, inkl. moms (fra Kanpla)", 7)
    ws.merge_cells(start_row=7, start_column=e0, end_row=7, end_column=e0 + len(eks) - 1)
    for i, k in enumerate(eks):
        c = e0 + i
        hoved(c, k["navn"])
        ws.cell(5, c, k.get("pris_pr_stk")).number_format = KR
        ws.cell(6, c, k["sats"]).number_format = "0%"
        for r in (5, 6):
            ws.cell(r, c).fill, ws.cell(r, c).border = GUL, RAMME
    ws.cell(5, c_eks, "Pris pr. stk (antal-kolonner)").font = Font(italic=True, size=9)
    ws.cell(6, c_eks, "Vareforbrugssats").font = Font(italic=True, size=9)

    kp = f"$C$4"
    moms = cfg["moms_faktor"]
    for r in range(F, S + 1):
        dag = date(aar, md, r - F + 1)
        ws.cell(r, 1, dag).number_format = "d"
        A = f"$A{r}"
        for c in range(lev0, lev_sidst + 1):
            ws.cell(r, c, f'=SUMIFS(Ordrer[Total],Ordrer[Leveringsdato],{A},Ordrer[Kolonne],{L(c)}$8)').fill = BLAA
        ws.cell(r, c_ialt, f"=SUM({L(lev0)}{r}:{L(lev_sidst)}{r})+{L(c_oevr)}{r}")
        ws.cell(r, c_mtd, f"={L(c_ialt)}{r}" + (f"+{L(c_mtd)}{r - 1}" if r > F else ""))
        ws.cell(r, c_oevr).fill = GUL
        ws.cell(r, c_kp, f"={kp}")
        ws.cell(r, c_kuv, f'=SUMIFS(Kanpla[kuverter],Kanpla[Dato],{A})').fill = BLAA
        ws.cell(r, c_kb, f"={L(c_kp)}{r}*{L(c_kuv)}{r}")
        kuv_cum = f"SUM({L(c_kuv)}${F}:{L(c_kuv)}{r})"
        ws.cell(r, c_kpk, f'=IF({kuv_cum}=0,"",({L(c_mtd)}{r}-{L(c_tilv)}{r})/{kuv_cum})')
        ws.cell(r, c_lb, f"=SUM({L(c_kb)}${F}:{L(c_kb)}{r})+{L(c_tilv)}{r}")
        ws.cell(r, c_diff, f"={L(c_lb)}{r}-{L(c_mtd)}{r}")
        led_inkl, led_vare = [], []
        for i, k in enumerate(eks):
            c = e0 + i
            ws.cell(r, c, f'=SUMIFS(Kanpla[{k["noegle"]}],Kanpla[Dato],{A})').fill = BLAA
            v = f"{L(c)}{r}" + (f"*{L(c)}$5" if k.get("pris_pr_stk") else "")
            led_inkl.append(v)
            led_vare.append(f"{v}*{L(c)}$6")
        ws.cell(r, c_eks, "=" + "+".join(led_inkl)).font = FED
        ws.cell(r, c_vare, f"=({'+'.join(led_vare)})/{moms}").font = FED
        ws.cell(r, c_tilv, f"={L(c_vare)}{r}" + (f"+{L(c_tilv)}{r - 1}" if r > F else ""))
        for c in range(1, c_tilv + 1):
            cell = ws.cell(r, c)
            if cell.value is not None or c in (c_oevr,):
                cell.border = RAMME
                if c > 1:
                    cell.number_format = "0" if c == c_kuv else KR
        if dag.weekday() >= 5:
            ws.cell(r, 1).font = Font(color="999999")

    # Total-række
    T = S + 1
    ws.cell(T, 1, "Total").font = FED
    for c in list(range(lev0, c_ialt + 1)) + [c_oevr, c_kuv, c_kb] + list(range(e0, c_vare + 1)):
        cell = ws.cell(T, c, f"=SUM({L(c)}{F}:{L(c)}{S})")
        cell.font, cell.border = FED, RAMME
        cell.number_format = "0" if c == c_kuv else KR
    for c in (c_kpk, c_lb, c_diff):
        cell = ws.cell(T, c, f"={L(c)}{S}")
        cell.font, cell.border, cell.number_format = FED, RAMME, KR

    # Grøn/rød på kuvertpris (faktisk kostpris pr. kuvert mod budget) og diff
    rng = f"{L(c_kpk)}{F}:{L(c_kpk)}{T}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThanOrEqual", formula=[kp],
                                                  fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=[kp],
                                                  fill=PatternFill("solid", fgColor="FFC7CE")))
    rng = f"{L(c_diff)}{F}:{L(c_diff)}{T}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"],
                                                  font=Font(color="C00000", bold=True)))

    ws.column_dimensions["A"].width = 11
    for c in range(2, c_tilv + 1):
        ws.column_dimensions[L(c)].width = 13
    for c in (c_mtd + 1, c_oevr + 1, c_diff + 1):
        ws.column_dimensions[L(c)].width = 2
    ws.row_dimensions[8].height = 45
    ws.freeze_panes = ws.cell(F, 2)


def byg_excel(sti, ordrer, kanpla, cfg, maaneder):
    wb = Workbook()
    wb.remove(wb.active)
    ordrer = sorted(ordrer, key=lambda o: (o["leveringsdato"], str(o["ordre"])))
    ordrer_ark(wb, ordrer)
    linje_ark(wb, ordrer)
    kanpla_kol = kanpla_ark(wb, kanpla, cfg)
    for aar, md in sorted(maaneder):
        maaneds_ark(wb, aar, md, cfg, kanpla_kol)
    wb.active = 0
    wb.save(sti)
