"""K-flow kommandoer.

  python koer.py skabelon --fra 2026-10 --til 2026-12     # ny tom K-flow.xlsx (lægges på SharePoint)
  python koer.py dashboard --excel <sti til K-flow.xlsx>    # læser arket -> web/data.js
  python koer.py importer-mails <mappe> --excel <sti>       # indlæs gemte mails (test / gamle mails)
  python koer.py importer-kanpla <mappe> --excel <sti>      # indlæs Kanpla-CSV
  python koer.py demo                                       # alt ovenstående med demo-data

I drift fylder Power Automate + Office Scriptet arket med mails; her skal kun 'dashboard' køres.
"""
import argparse
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from kkflow.arbejdsbog import importer_kanpla, importer_ordre, laes_data, laes_regler
from kkflow.beregning import beregn_maaned
from kkflow.excel_export import lav_skabelon
from kkflow.kanpla import laes_kanpla
from kkflow.parse_ordre import parse_ordrekvittering

ROD = Path(__file__).resolve().parent
CFG = json.loads((ROD / "config.json").read_text(encoding="utf-8"))
STANDARD_EXCEL = ROD / "output" / "K-flow.xlsx"


def maaneder_mellem(fra, til):
    y, m = map(int, fra.split("-"))
    ty, tm = map(int, til.split("-"))
    ud = []
    while (y, m) <= (ty, tm):
        ud.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return ud


def skabelon(sti, fra, til):
    sti.parent.mkdir(parents=True, exist_ok=True)
    lav_skabelon(sti, CFG, maaneder_mellem(fra, til))
    print(f"Skabelon: {sti}")


def importer_mails(mappe, sti):
    wb = load_workbook(sti)
    regler = laes_regler(wb)
    optalt = {}
    for fil in sorted(p for p in Path(mappe).iterdir() if p.suffix.lower() in (".eml", ".html", ".htm", ".txt")):
        o = parse_ordrekvittering(fil, regler)
        status = importer_ordre(wb, o)
        optalt[status] = optalt.get(status, 0) + 1
        if status not in ("indlæst", "dublet") or o["advarsler"]:
            print(f"  {fil.name}: {status} {'; '.join(o['advarsler'])}")
    wb.save(sti)
    print("Mails:", ", ".join(f"{v} {k}" for k, v in optalt.items()))


def importer_kanpla_cmd(mappe, sti):
    wb = load_workbook(sti)
    n = importer_kanpla(wb, laes_kanpla(mappe))
    wb.save(sti)
    print(f"Kanpla: {n} nye dage")


def dashboard(sti):
    wb = load_workbook(sti, read_only=False, data_only=True)
    ordrer, kanpla = laes_data(wb)
    maaneder = {tuple(map(int, d[:7].split("-"))) for d in [o["leveringsdato"] for o in ordrer if o["leveringsdato"]] + list(kanpla)}
    web = {
        "opdateret": datetime.now().isoformat(timespec="minutes"),
        "kilde": Path(sti).name,
        "kantine": CFG["kantine"],
        "kuvertpris": CFG["kuvertpris"],
        "kolonner": [k["navn"] for k in CFG["indkoebs_kolonner"]],
        "maaneder": {f"{y}-{m:02d}": beregn_maaned(y, m, ordrer, kanpla, CFG) for y, m in sorted(maaneder)},
        "ordrer": sorted(ordrer, key=lambda o: (o["leveringsdato"], o["ordre"]), reverse=True),
    }
    js = ROD / "web" / "data.js"
    js.write_text("window.KK_DATA = " + json.dumps(web, ensure_ascii=False) + ";\n", encoding="utf-8")
    print(f"Dashboard-data: {js} ({len(ordrer)} ordrer, {len(kanpla)} Kanpla-dage)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("skabelon")
    s.add_argument("--fra", required=True)
    s.add_argument("--til", required=True)
    s.add_argument("--excel", type=Path, default=STANDARD_EXCEL)
    for navn in ("importer-mails", "importer-kanpla"):
        s = sub.add_parser(navn)
        s.add_argument("mappe", type=Path)
        s.add_argument("--excel", type=Path, default=STANDARD_EXCEL)
    sub.add_parser("dashboard").add_argument("--excel", type=Path, default=STANDARD_EXCEL)
    sub.add_parser("demo")
    a = ap.parse_args()

    if a.cmd == "skabelon":
        skabelon(a.excel, a.fra, a.til)
    elif a.cmd == "importer-mails":
        importer_mails(a.mappe, a.excel)
    elif a.cmd == "importer-kanpla":
        importer_kanpla_cmd(a.mappe, a.excel)
    elif a.cmd == "dashboard":
        dashboard(a.excel)
    elif a.cmd == "demo":
        sti = ROD / "output" / "K-flow-demo.xlsx"
        skabelon(sti, "2026-10", "2026-12")
        importer_mails(ROD / "eksempler" / "indbakke", sti)
        importer_kanpla_cmd(ROD / "eksempler" / "kanpla", sti)
        dashboard(sti)


if __name__ == "__main__":
    main()
